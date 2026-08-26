from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Sequence

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

import ontology_aligner_runtime as core

from . import config, runtime


def json_cell(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def result_frame(records: Sequence[dict[str, Any]]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for record in records:
        candidate_ids = list(record.get("candidate_ids", []))
        ranking = list(record.get("ranking", candidate_ids))
        initial_ranking = list(record.get("initial_ranking", ranking))
        accepted = list(record.get("accepted_ids", []))
        row = {
            "experiment": record.get("experiment", ""),
            "condition": record.get("condition", ""),
            "backbone_key": record.get("backbone_key", ""),
            "embedding_model": record.get("embedding_model", ""),
            "candidate_k": record.get("candidate_k", len(candidate_ids)),
            "dataset_key": record["dataset_key"],
            "sample_id": record["sample_id"],
            "source_sample_id": int(record["source_sample_id"]),
            "source_excel_row": int(record["source_excel_row"]),
            "query": record["query"],
            "preferred_name": record.get("preferred_name", ""),
            "gold_id": record.get("gold_id", ""),
            "accepted_gold_ids_json": json_cell(accepted),
            "candidate_ids_json": json_cell(candidate_ids),
            "candidate_scores_json": json_cell(record.get("candidate_scores", [])),
            "candidate_matched_surfaces_json": json_cell(
                record.get("candidate_matched_surfaces", [])
            ),
            "retrieval_top1_id": candidate_ids[0] if candidate_ids else "",
            "retrieval_correct": bool(candidate_ids and candidate_ids[0] in set(accepted)),
            "lcr_ranking_json": json_cell(initial_ranking),
            "lcr_top1_id": initial_ranking[0] if initial_ranking else "No Match",
            "lcr_correct": bool(initial_ranking and initial_ranking[0] in set(accepted)),
            "final_ranking_json": json_cell(ranking),
            "final_top1_id": ranking[0] if ranking else record.get("predicted_id", ""),
            "final_correct": bool(record.get("correct", False)),
            "graph_triggered": bool(record.get("graph_triggered", False)),
            "graph_rerank_applied": bool(
                record.get("graph_rerank_applied", False)
            ),
            "graph_changed_top1": bool(
                record.get("graph_rerank_applied", False)
                and record.get("initial_predicted_id") != record.get("predicted_id")
            ),
            "hpo_graph_context_json": json_cell(
                record.get("hpo_graph_context", [])
            ),
            "llm_response_key": record.get("llm_response_key", ""),
            "llm_ranking_repair_json": json_cell(
                record.get("llm_ranking_repair", {})
            ),
            "hgr_response_key": record.get("hgr_response_key", ""),
            "hgr_ranking_repair_json": json_cell(
                record.get("hgr_ranking_repair", {})
            ),
            "llm_attempts": int(record.get("llm_attempts", 0)),
            "hgr_attempts": int(record.get("hgr_attempts", 0)),
            "prompt_tokens": int(record.get("prompt_tokens", 0)),
            "completion_tokens": int(record.get("completion_tokens", 0)),
            "total_tokens": int(record.get("total_tokens", 0)),
            "hgr_prompt_tokens": int(record.get("hgr_prompt_tokens", 0)),
            "hgr_completion_tokens": int(record.get("hgr_completion_tokens", 0)),
            "hgr_total_tokens": int(record.get("hgr_total_tokens", 0)),
        }
        if "raw_candidate_ids" in record:
            raw_ids = list(record["raw_candidate_ids"])
            row.update(
                {
                    "raw_candidate_ids_json": json_cell(raw_ids),
                    "raw_candidate_scores_json": json_cell(
                        record.get("raw_candidate_scores", [])
                    ),
                    "raw_gold_candidate_rank": next(
                        (
                            index
                            for index, ontology_id in enumerate(raw_ids, 1)
                            if ontology_id in set(accepted)
                        ),
                        None,
                    ),
                    "raw_top1_correct": bool(raw_ids and raw_ids[0] in set(accepted)),
                }
            )
        rows.append(row)
    return pd.DataFrame(rows)


def retrieval_metrics(records: Sequence[dict[str, Any]], key: str) -> dict[str, Any]:
    ranks: list[int | None] = []
    for record in records:
        candidates = list(record[key])
        accepted = set(record["accepted_ids"])
        ranks.append(
            next(
                (
                    index
                    for index, ontology_id in enumerate(candidates, 1)
                    if ontology_id in accepted
                ),
                None,
            )
        )
    return {
        "rows": len(records),
        "top1": float(np.mean([rank == 1 for rank in ranks])),
        "recall_at_1": float(np.mean([rank is not None and rank <= 1 for rank in ranks])),
        "recall_at_3": float(np.mean([rank is not None and rank <= 3 for rank in ranks])),
        "recall_at_5": float(np.mean([rank is not None and rank <= 5 for rank in ranks])),
        "recall_at_10": float(
            np.mean([rank is not None and rank <= 10 for rank in ranks])
        ),
        "recall_at_20": float(
            np.mean([rank is not None and rank <= 20 for rank in ranks])
        ),
        "mrr": float(np.mean([1.0 / rank if rank else 0.0 for rank in ranks])),
        "mean_gold_rank": float(np.mean([rank for rank in ranks if rank]))
        if any(ranks)
        else None,
    }


def stage_records(
    records: Sequence[dict[str, Any]], stage: str
) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for record in records:
        item = dict(record)
        if stage == "Retrieval":
            item["predicted_id"] = item["candidate_ids"][0]
        elif stage == "LCR":
            ranking = list(item.get("initial_ranking", item.get("ranking", [])))
            item["predicted_id"] = ranking[0] if ranking else "No Match"
        elif stage == "Final":
            item["predicted_id"] = item.get("predicted_id", item["candidate_ids"][0])
        else:
            raise ValueError(stage)
        item["correct"] = item["predicted_id"] in set(item["accepted_ids"])
        output.append(item)
    return output


def summary_frame(
    records: Sequence[dict[str, Any]], group_columns: Sequence[str]
) -> pd.DataFrame:
    grouped: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        grouped[tuple(record.get(column, "") for column in group_columns)].append(record)
    rows: list[dict[str, Any]] = []
    for group, values in grouped.items():
        group_values = dict(zip(group_columns, group))
        by_dataset = runtime.group_by_dataset(values)
        for dataset in config.DATASET_ORDER:
            dataset_values = by_dataset[dataset]
            if not dataset_values:
                continue
            retrieval = retrieval_metrics(dataset_values, "candidate_ids")
            raw = (
                retrieval_metrics(dataset_values, "raw_candidate_ids")
                if "raw_candidate_ids" in dataset_values[0]
                else None
            )
            lcr = stage_records(dataset_values, "LCR")
            final = stage_records(dataset_values, "Final")
            row = {
                **group_values,
                "dataset_key": dataset,
                **{f"retrieval_{key}": value for key, value in retrieval.items()},
                "lcr_top1": core.method_metrics(lcr)["micro_top1"],
                "final_top1": core.method_metrics(final)["micro_top1"],
                "hgr_triggered": sum(
                    bool(value.get("graph_rerank_applied")) for value in dataset_values
                ),
                "hgr_changed_top1": sum(
                    bool(value.get("graph_rerank_applied"))
                    and value.get("initial_predicted_id") != value.get("predicted_id")
                    for value in dataset_values
                ),
                "hgr_wrong_to_right": sum(
                    not before["correct"] and after["correct"]
                    for before, after in zip(lcr, final)
                ),
                "hgr_right_to_wrong": sum(
                    before["correct"] and not after["correct"]
                    for before, after in zip(lcr, final)
                ),
            }
            if raw is not None:
                row.update({f"raw_{key}": value for key, value in raw.items()})
            rows.append(row)
        retrieval = retrieval_metrics(values, "candidate_ids")
        raw = (
            retrieval_metrics(values, "raw_candidate_ids")
            if "raw_candidate_ids" in values[0]
            else None
        )
        lcr = stage_records(values, "LCR")
        final = stage_records(values, "Final")
        lcr_metrics = core.method_metrics(lcr)
        final_metrics = core.method_metrics(final)
        overall = {
            **group_values,
            "dataset_key": "OVERALL",
            **{f"retrieval_{key}": value for key, value in retrieval.items()},
            "retrieval_macro_top1": float(
                np.mean(
                    [
                        retrieval_metrics(dataset_values, "candidate_ids")["top1"]
                        for dataset_values in runtime.group_by_dataset(values).values()
                        if dataset_values
                    ]
                )
            ),
            "lcr_top1": lcr_metrics["micro_top1"],
            "lcr_macro_top1": lcr_metrics["macro_top1"],
            "final_top1": final_metrics["micro_top1"],
            "final_macro_top1": final_metrics["macro_top1"],
            "hgr_triggered": sum(
                bool(value.get("graph_rerank_applied")) for value in values
            ),
            "hgr_changed_top1": sum(
                bool(value.get("graph_rerank_applied"))
                and value.get("initial_predicted_id") != value.get("predicted_id")
                for value in values
            ),
            "hgr_wrong_to_right": sum(
                not before["correct"] and after["correct"]
                for before, after in zip(lcr, final)
            ),
            "hgr_right_to_wrong": sum(
                before["correct"] and not after["correct"]
                for before, after in zip(lcr, final)
            ),
        }
        if raw is not None:
            overall.update({f"raw_{key}": value for key, value in raw.items()})
            overall["raw_macro_top1"] = float(
                np.mean(
                    [
                        retrieval_metrics(dataset_values, "raw_candidate_ids")["top1"]
                        for dataset_values in runtime.group_by_dataset(values).values()
                        if dataset_values
                    ]
                )
            )
        rows.append(overall)
    return pd.DataFrame(rows)


def config_frame(values: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "parameter": key,
                "value": json.dumps(value, ensure_ascii=False)
                if isinstance(value, (dict, list, tuple))
                else value,
            }
            for key, value in values.items()
        ]
    )


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    for index, sheet in enumerate(workbook.worksheets, 1):
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="176567")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for column in sheet.columns:
            letter = column[0].column_letter
            maximum = max(len(str(cell.value or "")) for cell in column[:100])
            sheet.column_dimensions[letter].width = min(max(maximum + 2, 12), 55)
        if sheet.max_row >= 2 and sheet.max_column >= 1:
            name = re_table_name(f"{path.stem}_{sheet.title}_{index}")
            table = Table(displayName=name, ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
            )
            sheet.add_table(table)
    workbook.save(path)


def re_table_name(value: str) -> str:
    return "T" + "".join(character for character in value if character.isalnum())[:200]


def write_experiment_workbook(
    experiment: str,
    records: Sequence[dict[str, Any]],
    group_columns: Sequence[str],
    run_config: dict[str, Any],
    output_path: Path | None = None,
) -> Path:
    path = output_path or config.EXPERIMENT_FILES[experiment]
    path.parent.mkdir(parents=True, exist_ok=True)
    frame = result_frame(records)
    summary = summary_frame(records, group_columns)
    comparisons = run_config.get("paired_comparisons", {})
    for index, row in summary.iterrows():
        if row.get("dataset_key") != "OVERALL":
            continue
        lookup_values = [str(row.get(column, "")) for column in group_columns]
        lookup = "|".join(lookup_values)
        comparison = comparisons.get(lookup) or comparisons.get("OVERALL")
        if comparison:
            summary.loc[index, "comparison_label"] = comparison.get("comparison", "")
            summary.loc[index, "macro_delta"] = comparison.get("macro_delta")
            interval = comparison.get("bootstrap_95_ci", [None, None])
            summary.loc[index, "bootstrap_95_ci_low"] = interval[0]
            summary.loc[index, "bootstrap_95_ci_high"] = interval[1]
            summary.loc[index, "mcnemar_exact_p"] = comparison.get(
                "mcnemar_exact_p"
            )
            summary.loc[index, "paired_wrong_to_right"] = comparison.get(
                "wrong_to_right"
            )
            summary.loc[index, "paired_right_to_wrong"] = comparison.get(
                "right_to_wrong"
            )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Results", index=False)
        summary.to_excel(writer, sheet_name="Summary", index=False)
        config_frame(run_config).to_excel(writer, sheet_name="Config", index=False)
    style_workbook(path)
    return path


def generate_summary_workbook(experiment_config: dict[str, Any]) -> Path:
    missing = [path for path in config.EXPERIMENT_FILES.values() if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Cannot generate SUMMARY.xlsx; missing {missing}")
    main_frames = []
    for experiment in ("E1", "E2", "E3", "E4"):
        frame = pd.read_excel(config.EXPERIMENT_FILES[experiment], sheet_name="Summary")
        frame.insert(0, "experiment", experiment)
        main_frames.append(frame)
    e5 = pd.read_excel(config.EXPERIMENT_FILES["E5"], sheet_name="Summary")
    e6 = pd.read_excel(config.EXPERIMENT_FILES["E6"], sheet_name="Summary")
    path = config.RESULTS_DIR / "SUMMARY.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.concat(main_frames, ignore_index=True).to_excel(
            writer, sheet_name="Main_Ablation", index=False
        )
        e5.to_excel(writer, sheet_name="E5_Embeddings", index=False)
        e6.to_excel(writer, sheet_name="E6_Candidate_K", index=False)
        config_frame(experiment_config).to_excel(
            writer, sheet_name="Config", index=False
        )
    style_workbook(path)
    return path
