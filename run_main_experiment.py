from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any, Sequence

import chromadb
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

import ontology_aligner_runtime as core
from dataset_utils import (
    DATASET_CHOICES,
    HpoCanonicalizer,
    canonical_gold_ids,
    dataset_config,
    load_gold_labels,
)
from ontology_aligner_oar import load_default_oar_runtime


ROOT = Path(__file__).resolve().parent
RESULTS_ROOT = ROOT / "results"
PRECOMPUTE_RESULTS_DIR = RESULTS_ROOT / "precompute"
RERANK_RESULTS_DIR = RESULTS_ROOT / "rerank"
EMBEDDING_CACHE = PRECOMPUTE_RESULTS_DIR / "query_embeddings.sqlite3"
RETRIEVAL_TOP_K = 20
DATASET_ORDER = (
    "genereviews-10",
    "id-68",
    "gsc2017",
    "gsc2024",
    "csc",
    "fgdd",
    "bc8_t3",
)
NOTEBOOK_DATASET_KEYS = {
    "fgdd": "fgdd_phenotype",
    "genereviews-10": "genereviews_10",
    "id-68": "id_68",
    "gsc2017": "gsc2017",
    "gsc2024": "gsc2024",
    "csc": "csc",
    "bc8_t3": "bc8_t3",
}


def safe_tag(value: str, limit: int = 48) -> str:
    tag = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._-").lower()
    return (tag or "model")[:limit]


def notebook_dataset_key(dataset: str) -> str:
    return NOTEBOOK_DATASET_KEYS[dataset]


def embedding_tag() -> str:
    return safe_tag(f"{core.EMBEDDING_MODEL}_oar")


def precompute_workbook_path(dataset: str) -> Path:
    return PRECOMPUTE_RESULTS_DIR / (
        f"{notebook_dataset_key(dataset)}_retrieval_top{RETRIEVAL_TOP_K}_"
        f"{embedding_tag()}_dense-only.xlsx"
    )


def experiment_key() -> str:
    identity = experiment_identity()
    identity_hash = core.stable_hash(identity)[:12]
    return f"{safe_tag(str(identity['llm_model']))}_{identity_hash}"


def experiment_dir() -> Path:
    return RERANK_RESULTS_DIR / ".runs" / experiment_key()


def llm_cache_dir() -> Path:
    return experiment_dir() / "llm_cache"


def experiment_manifest_path() -> Path:
    return experiment_dir() / "experiment_manifest.json"


def dataset_dir(dataset: str) -> Path:
    return experiment_dir() / dataset


def prompt_hashes() -> dict[str, str]:
    return {
        "lcr_system": hashlib.sha256(core.SYSTEM_PROMPT.encode()).hexdigest(),
        "lcr_user_template": hashlib.sha256(
            core.USER_PROMPT_TEMPLATE.encode()
        ).hexdigest(),
        "hgr_system": hashlib.sha256(
            core.HPO_GRAPH_SYSTEM_PROMPT.encode()
        ).hexdigest(),
        "hgr_user_template": hashlib.sha256(
            core.HPO_GRAPH_RERANK_PROMPT.encode()
        ).hexdigest(),
    }


def experiment_identity() -> dict[str, Any]:
    config = core.load_config()
    runtime = load_default_oar_runtime(ROOT)
    return {
        "method": "OntologyAligner",
        "pipeline": ["OAR", "LCR", "HGR"],
        "llm_model": str(config["llm"]["model"]),
        "temperature": core.LLM_TEMPERATURE,
        "embedding_model": core.EMBEDDING_MODEL,
        "retrieval_top_k": RETRIEVAL_TOP_K,
        "candidate_count": core.LLM_CANDIDATE_COUNT,
        "hgr_probe_top_k": core.HPO_GRAPH_PROBE_TOP_K,
        "hgr_max_ancestor_distance": core.HPO_MAX_ANCESTOR_DISTANCE,
        "max_attempts": core.MAX_ATTEMPTS,
        "prompt_hashes": prompt_hashes(),
        "hpo_jsonl_sha256": core.file_sha256(core.HPO_JSONL_PATH),
        "hpo_graph_sha256": core.file_sha256(core.HPO_GRAPH_PATH),
        "oar_run_id": runtime.run_id,
        "oar_model_sha256": runtime.model_sha256,
    }


def ensure_experiment_manifest() -> dict[str, Any]:
    manifest_path = experiment_manifest_path()
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    identity = experiment_identity()
    if manifest_path.exists():
        existing = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
        if existing["identity"] != identity:
            raise RuntimeError(
                "The frozen main-experiment identity changed; use a new result directory"
            )
        return existing
    manifest = {
        "created_at": core.utc_now(),
        "identity": identity,
        "rerank_notebook_sha256": core.file_sha256(core.RERANK_NOTEBOOK_PATH),
        "datasets": list(DATASET_ORDER),
        "llm_cache_policy": "new main-experiment cache; no ablation response reuse",
        "deterministic_cache_policy": "reuse source embeddings and persistent OAR index",
        "no_match_policy": "valid LCR response; HGR must return a complete ranking",
        "rate_policy": [
            {"requests_per_minute": 600, "max_concurrency": 30},
            {"requests_per_minute": 300, "max_concurrency": 15},
            {"requests_per_minute": 150, "max_concurrency": 8},
        ],
    }
    core.write_json(manifest_path, manifest)
    return manifest


def load_samples(dataset: str) -> list[dict[str, Any]]:
    config = dataset_config(dataset)
    frame = pd.read_excel(config.workbook, sheet_name=config.sheet, dtype=str).fillna("")
    frame = frame.reset_index(drop=True)
    samples: list[dict[str, Any]] = []
    for row in frame.to_dict(orient="records"):
        row_index = len(samples)
        metadata = {
            column: str(row.get(column, ""))
            for column in config.inference_columns
            if column != "Raw_Phenotype_Names"
        }
        samples.append(
            {
                "sample_id": f"{dataset}:{row_index:06d}",
                "dataset_key": dataset,
                "source_sample_id": row_index,
                "source_excel_row": row_index + 2,
                "query": str(row["Raw_Phenotype_Names"]),
                "preferred_name": str(row.get("Standard_Phenotype_Names", "")),
                "gold_id": "",
                "accepted_ids": [],
                "metadata": metadata,
            }
        )
    return samples


def validate_oar_records(
    records: Sequence[dict[str, Any]],
    dataset: str,
    expected_rows: int,
) -> None:
    if len(records) != expected_rows:
        raise ValueError(f"{dataset}: OAR row count is {len(records)}, expected {expected_rows}")
    for record in records:
        candidate_ids = list(record.get("candidate_ids", []))
        candidate_scores = list(record.get("candidate_scores", []))
        matched_surfaces = list(record.get("candidate_matched_surfaces", []))
        if (
            len(candidate_ids) != RETRIEVAL_TOP_K
            or len(set(candidate_ids)) != RETRIEVAL_TOP_K
            or len(candidate_scores) != RETRIEVAL_TOP_K
            or len(matched_surfaces) != RETRIEVAL_TOP_K
        ):
            raise ValueError(
                f"{dataset}: sample {record.get('sample_id')} does not contain "
                f"{RETRIEVAL_TOP_K} unique OAR candidates"
            )


def normalize_oar_records(
    records: Sequence[dict[str, Any]],
    samples: Sequence[dict[str, Any]],
    dataset: str,
) -> list[dict[str, Any]]:
    validate_oar_records(records, dataset, len(samples))
    normalized: list[dict[str, Any]] = []
    for record, sample in zip(records, samples):
        if int(record["source_sample_id"]) != int(sample["source_sample_id"]):
            raise ValueError(f"{dataset}: OAR row order changed")
        if str(record["query"]) != str(sample["query"]):
            raise ValueError(f"{dataset}: OAR query changed at {sample['sample_id']}")
        candidate_ids = list(record["candidate_ids"])
        normalized.append(
            {
                **record,
                **sample,
                "method": "OAR",
                "candidate_ids": candidate_ids,
                "candidate_scores": list(record["candidate_scores"]),
                "candidate_matched_surfaces": list(
                    record["candidate_matched_surfaces"]
                ),
                "predicted_id": candidate_ids[0],
                "no_match": False,
                "correct": False,
            }
        )
    return normalized


def load_precompute_workbook(
    path: Path,
    dataset: str,
    samples: Sequence[dict[str, Any]],
) -> list[dict[str, Any]]:
    frame = pd.read_excel(path, sheet_name="Samples", dtype=str).fillna("")
    config_frame = pd.read_excel(path, sheet_name="Run_Config", dtype=str).fillna("")
    run_config = dict(zip(config_frame["parameter"], config_frame["value"]))
    required = {
        "sample_id",
        "source_excel_row",
        "query",
        "candidate_ids_json",
        "candidate_scores_json",
        "candidate_matched_surfaces_json",
    }
    missing = required - set(frame.columns)
    if missing:
        raise KeyError(f"{path} is missing precompute columns: {sorted(missing)}")
    expected_key = notebook_dataset_key(dataset)
    if run_config.get("dataset_key") != expected_key:
        raise ValueError(
            f"{path} dataset is {run_config.get('dataset_key')}, expected {expected_key}"
        )
    if int(float(run_config.get("retrieval_top_k", 0))) != RETRIEVAL_TOP_K:
        raise ValueError(f"{path} is not a Top-{RETRIEVAL_TOP_K} precompute workbook")
    if run_config.get("retrieval_method") != "oar_dense_cosine":
        raise ValueError(f"{path} does not contain OAR retrieval")
    if len(frame) != len(samples):
        raise ValueError(f"{path} has {len(frame)} rows, expected {len(samples)}")

    records: list[dict[str, Any]] = []
    for row, sample in zip(frame.to_dict(orient="records"), samples):
        source_sample_id = int(row["sample_id"])
        if source_sample_id != int(sample["source_sample_id"]):
            raise ValueError(f"{path} sample order changed at {source_sample_id}")
        records.append(
            {
                **sample,
                "method": "OAR",
                "candidate_ids": json.loads(row["candidate_ids_json"]),
                "candidate_scores": json.loads(row["candidate_scores_json"]),
                "candidate_matched_surfaces": json.loads(
                    row["candidate_matched_surfaces_json"]
                ),
                "predicted_id": json.loads(row["candidate_ids_json"])[0],
                "no_match": False,
                "correct": False,
            }
        )
    return normalize_oar_records(records, samples, dataset)


def style_sheet(workbook_path: Path, sheet_name: str, widths: dict[str, float]) -> None:
    workbook = load_workbook(workbook_path)
    sheet = workbook[sheet_name]
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="176567")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    if sheet.max_row >= 2 and sheet.max_column >= 1:
        table = Table(
            displayName=f"{sheet_name.replace('_', '')}Table", ref=sheet.dimensions
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    workbook.save(workbook_path)


def config_frame(run_config: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "parameter": key,
                "value": json.dumps(value, ensure_ascii=False)
                if isinstance(value, (dict, list))
                else value,
            }
            for key, value in run_config.items()
        ]
    )


def build_precompute_frame(
    records: Sequence[dict[str, Any]],
    terms: dict[str, Any],
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for record in sorted(records, key=lambda item: int(item["source_sample_id"])):
        accepted_ids = list(record["accepted_ids"])
        accepted_set = set(accepted_ids)
        gold_rank = next(
            (
                rank
                for rank, ontology_id in enumerate(record["candidate_ids"], 1)
                if ontology_id in accepted_set
            ),
            None,
        )
        row: dict[str, Any] = {
            "sample_id": int(record["source_sample_id"]),
            "source_excel_row": int(record["source_excel_row"]),
            "query": record["query"],
            "preferred_name": record["preferred_name"],
            "gold_id": record["gold_id"],
            "accepted_gold_ids_json": json.dumps(
                accepted_ids, ensure_ascii=False, separators=(",", ":")
            ),
            "gold_candidate_rank": gold_rank,
            "hit_at_1": bool(gold_rank and gold_rank <= 1),
            "candidate_ids_json": json.dumps(
                record["candidate_ids"], ensure_ascii=False, separators=(",", ":")
            ),
            "candidate_scores_json": json.dumps(
                record["candidate_scores"], separators=(",", ":")
            ),
            "candidate_matched_surfaces_json": json.dumps(
                record["candidate_matched_surfaces"],
                ensure_ascii=False,
                separators=(",", ":"),
            ),
        }
        row.update(record.get("metadata", {}))
        for rank, ontology_id in enumerate(record["candidate_ids"], 1):
            row[f"top_{rank:02d}_id"] = ontology_id
            row[f"top_{rank:02d}_name"] = terms[ontology_id].name
        rows.append(row)
    return pd.DataFrame(rows)


def export_precompute_workbook(
    dataset: str,
    records: Sequence[dict[str, Any]],
) -> Path:
    path = precompute_workbook_path(dataset)
    path.parent.mkdir(parents=True, exist_ok=True)
    terms = core.load_ontology()
    frame = build_precompute_frame(records, terms)
    runtime = load_default_oar_runtime(ROOT)
    dataset_spec = dataset_config(dataset)
    run_config = {
        "dataset_key": notebook_dataset_key(dataset),
        "dataset_path": str(dataset_spec.workbook.resolve()),
        "dataset_sheet": dataset_spec.sheet,
        "accepted_gold_policy": "current ontology ID plus official accepted historical IDs",
        "sample_count": len(frame),
        "ontology": "hpo",
        "ontology_release": "2026-06-23",
        "ontology_sha256": core.file_sha256(core.HPO_JSONL_PATH),
        "embedding_backend": "openai",
        "embedding_model": core.EMBEDDING_MODEL,
        "embedding_dimensions": 3072,
        "retrieval_top_k": RETRIEVAL_TOP_K,
        "retrieval_method": "oar_dense_cosine",
        "oar_enabled": True,
        "oar_run_id": runtime.run_id,
        "oar_model_sha256": runtime.model_sha256,
        "dense_collection_name": core.OAR_COLLECTION_NAME,
        "source_runner": str(Path(__file__).name),
        "created_at": core.utc_now(),
        "recall_at_1": float(frame["hit_at_1"].mean()),
    }
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Samples", index=False)
        config_frame(run_config).to_excel(writer, sheet_name="Run_Config", index=False)
    style_sheet(
        path,
        "Samples",
        {"A": 12, "B": 16, "C": 42, "D": 32, "E": 18, "F": 60, "G": 18},
    )
    style_sheet(path, "Run_Config", {"A": 36, "B": 100})
    return path


def prepare_oar(dataset: str, samples: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    workbook_path = precompute_workbook_path(dataset)
    if workbook_path.exists():
        records = load_precompute_workbook(workbook_path, dataset, samples)
        print(f"{dataset}: reuse {len(records)} OAR records from {workbook_path}", flush=True)
        return evaluate_records(dataset, records)

    config = core.load_config()
    resolver = core.EmbeddingResolver(config, EMBEDDING_CACHE)
    vectors, embedding_stats = resolver.resolve([sample["query"] for sample in samples])
    resolver.local.close()
    if resolver.source is not None:
        resolver.source.close()

    runtime = load_default_oar_runtime(ROOT)
    projected = runtime.project(vectors)
    client = chromadb.PersistentClient(path=str(core.OAR_CHROMA_PATH))
    collection = client.get_collection(core.OAR_COLLECTION_NAME)
    candidates = core.retrieve_many(collection, projected, top_k=RETRIEVAL_TOP_K)
    records = core.retrieval_records(samples, candidates, "OAR")
    validate_oar_records(records, dataset, len(samples))
    evaluated = evaluate_records(dataset, records)
    export_precompute_workbook(dataset, evaluated)
    output_dir = dataset_dir(dataset)
    core.write_jsonl(output_dir / "oar_records.jsonl", evaluated)
    core.write_json(
        output_dir / "oar_manifest.json",
        {
            "completed_at": core.utc_now(),
            "dataset": dataset,
            "dataset_sha256": core.file_sha256(dataset_config(dataset).workbook),
            "samples": len(evaluated),
            "embedding_stats": embedding_stats,
            "oar_run_id": runtime.run_id,
            "oar_model_sha256": runtime.model_sha256,
            "collection": core.OAR_COLLECTION_NAME,
        },
    )
    return evaluated


def run_lcr(
    dataset: str,
    oar_records: Sequence[dict[str, Any]],
    max_concurrency: int,
    requests_per_minute: int,
) -> list[dict[str, Any]]:
    terms = core.load_ontology()
    config = core.load_config()
    llm_config = config["llm"]
    model = str(llm_config["model"])
    runner = core.LLMTaskRunner(llm_config, model, requests_per_minute)
    tasks = [
        core.build_lcr_task(record, terms, "MAIN_LCR", model)
        for record in oar_records
    ]
    responses = core.run_unique_tasks(
        tasks,
        llm_cache_dir() / "lcr_responses.jsonl",
        runner,
        max_concurrency,
    )
    records = core.assemble_lcr_records(oar_records, tasks, responses, "LCR")
    core.write_jsonl(dataset_dir(dataset) / "lcr_records.jsonl", records)
    return records


def run_hgr(
    dataset: str,
    lcr_records: Sequence[dict[str, Any]],
    max_concurrency: int,
    requests_per_minute: int,
) -> list[dict[str, Any]]:
    terms = core.load_ontology()
    graph = core.load_hpo_graph()
    config = core.load_config()
    llm_config = config["llm"]
    model = str(llm_config["model"])
    runner = core.LLMTaskRunner(llm_config, model, requests_per_minute)

    triggered: list[dict[str, Any]] = []
    for record in lcr_records:
        context = core.build_hpo_graph_context(
            record["ranking"], record["candidate_ids"][0], terms, graph
        )
        record["hpo_graph_context"] = context
        record["graph_triggered"] = bool(context)
        if context:
            triggered.append(record)
    tasks = [
        core.build_hgr_task(record, terms, model, "MAIN_HGR")
        for record in triggered
    ]
    responses = core.run_unique_tasks(
        tasks,
        llm_cache_dir() / "hgr_responses.jsonl",
        runner,
        max_concurrency,
    )
    response_by_sample = {
        record["sample_id"]: (task, responses[task["key"]])
        for record, task in zip(triggered, tasks)
    }

    final_records: list[dict[str, Any]] = []
    for record in lcr_records:
        output = dict(record)
        output["method"] = "OntologyAligner"
        output["initial_ranking"] = list(record["ranking"])
        output["initial_predicted_id"] = record["predicted_id"]
        output["graph_rerank_applied"] = record["sample_id"] in response_by_sample
        if output["graph_rerank_applied"]:
            task, response = response_by_sample[record["sample_id"]]
            ranking = list(response["ranking"])
            if response["no_match"] or not ranking:
                raise ValueError("Main-experiment HGR produced No Match")
            output["ranking"] = ranking
            output["no_match"] = False
            output["predicted_id"] = ranking[0]
            output["hgr_response_key"] = task["key"]
            output["hgr_attempts"] = response["attempts"]
            output["hgr_api_seconds"] = response["api_seconds"]
            output["hgr_prompt_tokens"] = response["prompt_tokens"]
            output["hgr_completion_tokens"] = response["completion_tokens"]
            output["hgr_total_tokens"] = response["total_tokens"]
            output["hgr_raw_response"] = response["raw_response"]
        else:
            output["hgr_response_key"] = ""
            output["hgr_attempts"] = 0
            output["hgr_api_seconds"] = 0.0
            output["hgr_prompt_tokens"] = 0
            output["hgr_completion_tokens"] = 0
            output["hgr_total_tokens"] = 0
            output["hgr_raw_response"] = ""
        final_records.append(output)
    core.write_jsonl(dataset_dir(dataset) / "final_records.jsonl", final_records)
    return final_records


def evaluate_records(
    dataset: str,
    records: Sequence[dict[str, Any]],
) -> list[dict[str, Any]]:
    gold = load_gold_labels(dataset=dataset)
    if len(records) != len(gold):
        raise ValueError(f"{dataset}: prediction/gold row mismatch")
    canonicalizer = HpoCanonicalizer()
    evaluated: list[dict[str, Any]] = []
    for record, gold_row in zip(records, gold.to_dict(orient="records")):
        row_index = int(gold_row["row_index"])
        if int(record["source_sample_id"]) != row_index:
            raise ValueError(f"{dataset}: row order changed at {row_index}")
        accepted_ids = list(
            canonical_gold_ids(
                gold_row["HPO_id"],
                canonicalizer,
                gold_row["Accepted_HPO_ids"],
            )
        )
        output = dict(record)
        output["gold_id"] = str(gold_row["HPO_id"])
        output["accepted_ids"] = accepted_ids
        output["correct"] = output["predicted_id"] in set(accepted_ids)
        evaluated.append(output)
    return evaluated


def method_metrics(records: Sequence[dict[str, Any]]) -> dict[str, Any]:
    correct = sum(bool(record["correct"]) for record in records)
    return {
        "rows": len(records),
        "correct": correct,
        "top1_accuracy": correct / len(records),
        "no_match": sum(bool(record.get("no_match")) for record in records),
    }


def build_rerank_frame(records: Sequence[dict[str, Any]], terms: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for record in sorted(records, key=lambda item: int(item["source_sample_id"])):
        ranking = list(record.get("ranking", []))
        initial_ranking = list(record.get("initial_ranking", ranking))
        graph_context = list(record.get("hpo_graph_context", []))
        accepted_ids = list(record["accepted_ids"])
        accepted_set = set(accepted_ids)
        gold_rank = next(
            (
                rank
                for rank, ontology_id in enumerate(ranking, 1)
                if ontology_id in accepted_set
            ),
            None,
        )
        initial_attempts = int(record.get("llm_attempts", 0))
        graph_attempts = int(record.get("hgr_attempts", 0))
        initial_api_seconds = float(record.get("llm_api_seconds", 0.0))
        graph_api_seconds = float(record.get("hgr_api_seconds", 0.0))
        initial_prompt_tokens = int(record.get("prompt_tokens", 0))
        graph_prompt_tokens = int(record.get("hgr_prompt_tokens", 0))
        initial_completion_tokens = int(record.get("completion_tokens", 0))
        graph_completion_tokens = int(record.get("hgr_completion_tokens", 0))
        initial_total_tokens = int(record.get("total_tokens", 0))
        graph_total_tokens = int(record.get("hgr_total_tokens", 0))
        row: dict[str, Any] = {
            "sample_id": int(record["source_sample_id"]),
            "source_excel_row": int(record["source_excel_row"]),
            "query": record["query"],
            "preferred_name": record["preferred_name"],
            "gold_id": record["gold_id"],
            "accepted_gold_ids_json": json.dumps(
                accepted_ids, ensure_ascii=False, separators=(",", ":")
            ),
            "prediction_status": "no_match" if record.get("no_match") else "ranked",
            "graph_rerank_applied": bool(record.get("graph_rerank_applied", False)),
            "graph_context_count": len(graph_context),
            "initial_top1_id": initial_ranking[0] if initial_ranking else "No Match",
            "initial_top1_name": terms[initial_ranking[0]].name
            if initial_ranking
            else "",
            "predicted_top1_id": ranking[0] if ranking else "No Match",
            "predicted_top1_name": terms[ranking[0]].name if ranking else "",
            "gold_llm_rank": gold_rank,
            "hit_at_1": bool(gold_rank and gold_rank <= 1),
            "ranked_ids_json": json.dumps(
                ranking, ensure_ascii=False, separators=(",", ":")
            ),
            "initial_ranked_ids_json": json.dumps(
                initial_ranking, ensure_ascii=False, separators=(",", ":")
            ),
            "hpo_graph_context_json": json.dumps(
                graph_context, ensure_ascii=False, separators=(",", ":")
            ),
            "initial_attempts": initial_attempts,
            "graph_attempts": graph_attempts,
            "attempts": initial_attempts + graph_attempts,
            "initial_api_seconds": initial_api_seconds,
            "graph_api_seconds": graph_api_seconds,
            "api_seconds": initial_api_seconds + graph_api_seconds,
            "initial_prompt_tokens": initial_prompt_tokens,
            "graph_prompt_tokens": graph_prompt_tokens,
            "prompt_tokens": initial_prompt_tokens + graph_prompt_tokens,
            "initial_completion_tokens": initial_completion_tokens,
            "graph_completion_tokens": graph_completion_tokens,
            "completion_tokens": initial_completion_tokens + graph_completion_tokens,
            "initial_total_tokens": initial_total_tokens,
            "graph_total_tokens": graph_total_tokens,
            "total_tokens": initial_total_tokens + graph_total_tokens,
            "error": record.get("error", ""),
        }
        for rank, ontology_id in enumerate(ranking, 1):
            row[f"rank_{rank:02d}_id"] = ontology_id
            row[f"rank_{rank:02d}_name"] = terms[ontology_id].name
        rows.append(row)
    return pd.DataFrame(rows)


def rerank_workbook_path(dataset: str) -> Path:
    model = str(core.load_config()["llm"]["model"])
    return RERANK_RESULTS_DIR / (
        f"{notebook_dataset_key(dataset)}_llm_rerank_candidates{core.LLM_CANDIDATE_COUNT}_"
        f"{embedding_tag()}_{safe_tag(model)}_hpo-graph.xlsx"
    )


def export_rerank_workbook(
    dataset: str,
    records: Sequence[dict[str, Any]],
    max_concurrency: int,
    requests_per_minute: int,
) -> Path:
    path = rerank_workbook_path(dataset)
    path.parent.mkdir(parents=True, exist_ok=True)
    terms = core.load_ontology()
    frame = build_rerank_frame(records, terms)
    precompute_path = precompute_workbook_path(dataset)
    graph_applied = frame["graph_rerank_applied"].astype(bool)
    changed_top1 = graph_applied & (
        frame["initial_top1_id"] != frame["predicted_top1_id"]
    )
    identity = experiment_identity()
    run_config = {
        "dataset_key": notebook_dataset_key(dataset),
        "sample_count": len(frame),
        "ontology": "hpo",
        "accepted_gold_policy": "candidate hits any ID in accepted_gold_ids_json",
        "precompute_path": str(precompute_path.resolve()),
        "precompute_sha256": core.file_sha256(precompute_path),
        "embedding_model": core.EMBEDDING_MODEL,
        "oar_enabled": True,
        "llm_model": identity["llm_model"],
        "llm_temperature": core.LLM_TEMPERATURE,
        "llm_candidate_count": core.LLM_CANDIDATE_COUNT,
        "hpo_graph_rerank_enabled": True,
        "hpo_graph_path": str(core.HPO_GRAPH_PATH.resolve()),
        "hpo_graph_sha256": core.file_sha256(core.HPO_GRAPH_PATH),
        "hpo_graph_probe_top_k": core.HPO_GRAPH_PROBE_TOP_K,
        "hpo_max_ancestor_distance": core.HPO_MAX_ANCESTOR_DISTANCE,
        "hpo_graph_trigger_rule": "dense_top1_disagrees_and_is_related_to_llm_top3",
        "graph_rerank_sample_count": int(graph_applied.sum()),
        "graph_changed_top1_count": int(changed_top1.sum()),
        "no_match_count": int((frame["prediction_status"] == "no_match").sum()),
        "rerank_config_hash": experiment_key().rsplit("_", 1)[-1],
        "rerank_config_identity": identity,
        "llm_max_concurrency": max_concurrency,
        "llm_requests_per_minute": requests_per_minute,
        "max_attempts": core.MAX_ATTEMPTS,
        "checkpoint_path": str(dataset_dir(dataset).resolve()),
        "embedding_api_calls_this_runner": 0,
        "created_at": core.utc_now(),
        "accuracy_at_1": float(frame["hit_at_1"].mean()),
    }
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Rankings", index=False)
        config_frame(run_config).to_excel(writer, sheet_name="Run_Config", index=False)
    style_sheet(
        path,
        "Rankings",
        {
            "A": 12,
            "B": 16,
            "C": 42,
            "D": 30,
            "E": 18,
            "F": 60,
            "G": 18,
            "H": 18,
            "I": 18,
            "J": 18,
            "K": 30,
            "L": 18,
            "M": 30,
        },
    )
    style_sheet(path, "Run_Config", {"A": 36, "B": 100})
    return path


def write_dataset_workbook(
    dataset: str,
    oar_records: Sequence[dict[str, Any]],
    lcr_records: Sequence[dict[str, Any]],
    final_records: Sequence[dict[str, Any]],
) -> None:
    rows = []
    for oar, lcr, final in zip(oar_records, lcr_records, final_records):
        rows.append(
            {
                "row_index": final["source_sample_id"],
                "Raw_Phenotype_Names": final["query"],
                "HPO_id": final["gold_id"],
                "Accepted_HPO_ids": json.dumps(
                    final["accepted_ids"], ensure_ascii=False
                ),
                "OAR_top1": oar["predicted_id"],
                "LCR_top1": lcr["predicted_id"],
                "OntologyAligner_top1": final["predicted_id"],
                "HGR_triggered": final["graph_rerank_applied"],
                "correct": final["correct"],
            }
        )
    path = dataset_dir(dataset) / "OntologyAligner_predictions.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_excel(path, index=False, sheet_name="Predictions")


def finalize_dataset(
    dataset: str,
    oar_records: Sequence[dict[str, Any]],
    lcr_records: Sequence[dict[str, Any]],
    final_records: Sequence[dict[str, Any]],
    max_concurrency: int,
    requests_per_minute: int,
) -> dict[str, Any]:
    evaluated_oar = evaluate_records(dataset, oar_records)
    evaluated_lcr = evaluate_records(dataset, lcr_records)
    evaluated_final = evaluate_records(dataset, final_records)
    core.write_jsonl(dataset_dir(dataset) / "oar_records.jsonl", evaluated_oar)
    core.write_jsonl(dataset_dir(dataset) / "lcr_records.jsonl", evaluated_lcr)
    core.write_jsonl(dataset_dir(dataset) / "final_records.jsonl", evaluated_final)
    write_dataset_workbook(dataset, evaluated_oar, evaluated_lcr, evaluated_final)
    rerank_path = export_rerank_workbook(
        dataset, evaluated_final, max_concurrency, requests_per_minute
    )

    triggered = [
        record for record in evaluated_final if record["graph_rerank_applied"]
    ]
    lcr_by_id = {record["sample_id"]: record for record in evaluated_lcr}
    lcr_keys = {str(record["llm_response_key"]) for record in evaluated_lcr}
    hgr_keys = {
        str(record["hgr_response_key"])
        for record in triggered
        if record["hgr_response_key"]
    }
    summary = {
        "completed_at": core.utc_now(),
        "dataset": dataset,
        "dataset_sha256": core.file_sha256(dataset_config(dataset).workbook),
        "methods": {
            "OAR": method_metrics(evaluated_oar),
            "LCR": method_metrics(evaluated_lcr),
            "OntologyAligner": method_metrics(evaluated_final),
        },
        "hgr": {
            "triggered_samples": len(triggered),
            "trigger_rate": len(triggered) / len(evaluated_final),
            "changed_top1": sum(
                record["predicted_id"] != record["initial_predicted_id"]
                for record in triggered
            ),
            "wrong_to_right": sum(
                not lcr_by_id[record["sample_id"]]["correct"] and record["correct"]
                for record in triggered
            ),
            "right_to_wrong": sum(
                lcr_by_id[record["sample_id"]]["correct"] and not record["correct"]
                for record in triggered
            ),
        },
        "api_usage": {
            "LCR": core.response_usage_for_keys(
                llm_cache_dir() / "lcr_responses.jsonl", lcr_keys
            ),
            "HGR": core.response_usage_for_keys(
                llm_cache_dir() / "hgr_responses.jsonl", hgr_keys
            ),
        },
        "run_settings": {
            "max_concurrency": max_concurrency,
            "requests_per_minute": requests_per_minute,
        },
        "prompt_hashes": prompt_hashes(),
        "precompute_workbook": str(precompute_workbook_path(dataset)),
        "rerank_workbook": str(rerank_path),
    }
    core.write_json(dataset_dir(dataset) / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)
    return summary


def run_dataset(
    dataset: str,
    max_concurrency: int,
    requests_per_minute: int,
) -> dict[str, Any]:
    ensure_experiment_manifest()
    samples = load_samples(dataset)
    print(f"{dataset}: samples={len(samples)}", flush=True)
    oar_records = prepare_oar(dataset, samples)
    lcr_records = run_lcr(
        dataset, oar_records, max_concurrency, requests_per_minute
    )
    final_records = run_hgr(
        dataset, lcr_records, max_concurrency, requests_per_minute
    )
    return finalize_dataset(
        dataset,
        oar_records,
        lcr_records,
        final_records,
        max_concurrency,
        requests_per_minute,
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Run the frozen OntologyAligner main experiment by dataset"
    )
    parser.add_argument(
        "--dataset", choices=(*DATASET_CHOICES, "all"), required=True
    )
    parser.add_argument("--max-concurrency", type=int, default=30)
    parser.add_argument("--requests-per-minute", type=int, default=600)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    datasets = DATASET_ORDER if args.dataset == "all" else (args.dataset,)
    for dataset in datasets:
        run_dataset(dataset, args.max_concurrency, args.requests_per_minute)


if __name__ == "__main__":
    main()
